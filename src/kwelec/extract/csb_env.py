"""CSB Environment Statistics, electricity chapter (.xlsx), one workbook per year.

Column identity comes from config/editions/csb-env-<year>.yaml, never from the
header at run time: the extractor only checks that the header cells still say what
the reviewed config says they say, and stops if not. `draft()` writes a first
version of that config from the headers and the plant registry, for review."""

import calendar
import re
import warnings
from pathlib import Path

import openpyxl
import yaml

from ..normalize import clean, parse_printed, stored_number
from ..records import Obs
from .xlsx_fmt import displayed

warnings.filterwarnings("ignore", message="Sparkline Group extension")

MONTHS = {
    "يناير": 1, "فبراير": 2, "مارس": 3, "ابريل": 4, "أبريل": 4, "إبريل": 4, "مايو": 5, "يونيو": 6,
    "يوليو": 7, "اغسطس": 8, "أغسطس": 8, "سبتمبر": 9, "اكتوبر": 10, "أكتوبر": 10, "نوفمبر": 11, "ديسمبر": 12,
}
TOTAL_AR = {"الإجمالي", "الاجمالي", "الأجمالي", "المجموع"}

# English title -> (variable, fuel, unit as printed)
TABLES = [
    (r"^Electric Power Generated", "generation_gross", None, "MW/H"),
    (r"^Electrical Energy Consumed within", "station_use", None, "MW/H"),
    (r"^Exporting Electric Power", "sent_out", None, "MW/H"),
    (r"Natural Gas", "fuel_use", "natural_gas", "T.C.F"),
    (r"Gas Oil", "fuel_use", "gas_oil", "Barrel"),
    (r"Crude Oil", "fuel_use", "crude_oil", "Barrel"),
    (r"Heavy Oil", "fuel_use", "heavy_oil", "Barrel"),
    (r"^Information About Electric Power Stations", "plant_info", None, None),
]


def _s(v) -> str:
    return clean(str(v)) if v is not None else ""


def _key(s: str) -> str:
    return re.sub(r"[\s*]+", " ", s).strip().lower()


def load_aliases(registry: Path) -> dict[str, str]:
    plants = yaml.safe_load((registry / "plants.yaml").read_text())["plants"]
    return {_key(a): p["id"] for p in plants for a in p["aliases"] + [p["name_en"], p["name_ar"]]}


def _titles(ws):
    """(table number, Arabic title, English title, year) from the top of the sheet."""
    top = [(_s(ws.cell(r, c).value), ws.cell(r, c).value) for r in range(1, 6) for c in range(1, 5)]
    no = next((re.search(r"\(\s*(\d+)\s*\)", s).group(1) for s, _ in top if re.search(r"Table|جدول", s)), None)
    ar = next((s for s, _ in top if re.search(r"[؀-ۿ]", s) and "جدول" not in s and len(s) > 15), "")
    en = next((s for s, _ in top if re.match(r"[A-Za-z]", s) and "Table" not in s), "")
    year = next((v for _, v in top if isinstance(v, int) and 2000 < v < 2100), None)
    return no, ar, en, year


def _header_texts(ws, col: int, r0: int, r1: int) -> list[str]:
    return [_s(ws.cell(r, col).value) for r in range(r0, r1 + 1) if _s(ws.cell(r, col).value)]


def draft(xlsx: Path, source_id: str, year: int, registry: Path) -> dict:
    aliases = load_aliases(registry)
    wb = openpyxl.load_workbook(xlsx)
    tables = []
    for ws in wb.worksheets:
        no, ar, en, y = _titles(ws)
        kind = next((t for t in TABLES if re.search(t[0], en)), None)
        if not no or not kind or re.search(r"Water|Revenue|Cost", en):
            continue
        t = {"table_no": no, "sheet": ws.title, "title_en": en, "title_ar": ar, "year_printed": y}
        if kind[1] == "plant_info":
            t["layout"] = "plant_list"
            t["rows"] = []
            for r in range(5, ws.max_row + 1):
                # Each plant is an Arabic row carrying the values, then its English name below.
                pid = aliases.get(_key(_s(ws.cell(r, 2).value)))
                if pid and pid == aliases.get(_key(_s(ws.cell(r + 1, 2).value))) and ws.cell(r, 3).value is not None:
                    t["rows"].append({"row": r, "plant_id": pid,
                                      "labels": [_s(ws.cell(r, 2).value), _s(ws.cell(r + 1, 2).value)]})
            hdr = {c: _header_texts(ws, c, 5, 6) for c in range(3, 7)}
            t["columns"] = {}
            for c, texts in hdr.items():
                joined = " ".join(texts)
                var = ("operating_year" if "Operating Year" in joined else
                       "installed_capacity" if "Capacity" in joined else
                       "notes" if "Notes" in joined else None)
                if var:
                    t["columns"][openpyxl.utils.get_column_letter(c)] = {"variable": var, "header": texts,
                                                                         "unit_printed": {"installed_capacity": "MW", "operating_year": "year"}.get(var)}
            tables.append(t)
            continue
        # Monthly grid: find the month label column and the first data row.
        first = next(((r, c) for r in range(5, 15) for c in (2, 3) if _s(ws.cell(r, c).value) == "يناير"), None)
        if not first:
            raise ValueError(f"{source_id} {ws.title}: no January row")
        r_data, label_col = first
        t.update({"layout": "monthly", "variable": kind[1], "fuel": kind[2], "unit_printed": kind[3],
                  "label_col": openpyxl.utils.get_column_letter(label_col), "header_rows": [5, r_data - 1]})
        cols = {}
        for c in range(label_col + 1, ws.max_column + 1):
            texts = _header_texts(ws, c, 5, r_data - 1)
            if not texts:
                continue
            if texts == ["Month"]:  # English month labels repeated on the right (2013-2015)
                continue
            ids = {aliases[_key(x)] for x in texts if _key(x) in aliases}
            entry = {"header": texts}
            if any(re.search(r"^(Total|Net Export)", x) for x in texts):
                entry.update({"level": "national", "plant_id": None})
            elif len(ids) == 1:
                entry.update({"level": "plant", "plant_id": ids.pop()})
            else:
                entry.update({"level": "UNMAPPED", "plant_id": None})
            if kind[1] == "sent_out" and entry.get("plant_id", "") and entry["plant_id"].endswith(("-ro", "-boilers")):
                entry["variable"] = "desalination_use"
            if kind[1] == "sent_out" and "Distilled Water Production Plant" in texts[0]:
                entry["unit_group"] = "Shuwaikh distilled water production plant"
            if kind[1] == "sent_out" and entry["level"] == "national" and any("Net" in x for x in texts):
                entry["variable"] = "sent_out_net_total"
            cols[openpyxl.utils.get_column_letter(c)] = entry
        t["columns"] = cols
        tables.append(t)
    return {"source_id": source_id, "file": xlsx.name, "data_year": year, "tables": tables}


def _month_period(year: int, m: int) -> tuple[str, str]:
    return f"{year}-{m:02d}-01", f"{year}-{m:02d}-{calendar.monthrange(year, m)[1]:02d}"


def extract(xlsx: Path, cfg: dict) -> list[Obs]:
    wb_v = openpyxl.load_workbook(xlsx, data_only=True)   # cached values
    wb_f = openpyxl.load_workbook(xlsx)                   # formulas and formats
    year = cfg["data_year"]
    out: list[Obs] = []
    for t in cfg["tables"]:
        ws, wf = wb_v[t["sheet"]], wb_f[t["sheet"]]
        no, ar, en, y = _titles(wf)
        if (no, en, y) != (t["table_no"], t["title_en"], t["year_printed"]):
            raise ValueError(f"{cfg['source_id']} {t['sheet']}: title changed: {(no, en, y)}")
        if y != year:
            raise ValueError(f"{cfg['source_id']} table {no}: printed year {y} is not data year {year}")

        def cell(row: int, col_letter: str):
            c, f = ws[f"{col_letter}{row}"], wf[f"{col_letter}{row}"]
            formula = f.value if isinstance(f.value, str) and f.value.startswith("=") else None
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return str(stored_number(v)), displayed(v, f.number_format), "number", formula
            val, status = parse_printed(_s(v))
            return (str(val) if val is not None else None), _s(v), status, formula

        if t["layout"] == "plant_list":
            for r in t["rows"]:
                got = [_s(ws.cell(r["row"], 2).value), _s(ws.cell(r["row"] + 1, 2).value)]
                if got != r["labels"]:
                    raise ValueError(f"{cfg['source_id']} table {no} row {r['row']}: labels {got} != {r['labels']}")
                note = _s(ws[f"{next((k for k, c in t['columns'].items() if c['variable'] == 'notes'), 'F')}{r['row']}"].value)
                for col, c in t["columns"].items():
                    if c["variable"] == "notes":
                        continue
                    if _header_texts(wf, openpyxl.utils.column_index_from_string(col), 5, 6) != c["header"]:
                        raise ValueError(f"{cfg['source_id']} table {no} column {col}: header changed")
                    value, printed, status, formula = cell(r["row"], col)
                    out.append(Obs(cfg["source_id"], no, c["variable"], "plant", f"{year}-01-01", f"{year}-12-31",
                                   "annual", value, printed, status, c["unit_printed"], " / ".join(r["labels"]),
                                   " / ".join(c["header"]), "xlsx", plant_id=r["plant_id"], sheet=t["sheet"],
                                   cell_ref=f"{col}{r['row']}", cell_formula=formula,
                                   footnote_printed=note if note not in ("", "-") else None))
            continue

        label_col = t["label_col"]
        rows = []
        for r in range(t["header_rows"][1] + 1, ws.max_row + 1):
            lab = _s(ws[f"{label_col}{r}"].value)
            left = _s(ws[f"B{r}"].value)  # the total's label sits in the quarter column in some editions
            if lab in MONTHS:
                rows.append((r, lab, MONTHS[lab]))
            elif lab in TOTAL_AR or left in TOTAL_AR:
                lab = lab if lab in TOTAL_AR else left
                rows.append((r, lab, None))
                break
        if [m for _, _, m in rows] != list(range(1, 13)) + [None]:
            raise ValueError(f"{cfg['source_id']} table {no}: month rows {[m for _, _, m in rows]}")
        for col, c in t["columns"].items():
            if c["level"] == "UNMAPPED":
                raise ValueError(f"{cfg['source_id']} table {no} column {col}: unmapped {c['header']}")
            got = _header_texts(wf, openpyxl.utils.column_index_from_string(col), *t["header_rows"])
            if got != c["header"]:
                raise ValueError(f"{cfg['source_id']} table {no} column {col}: header {got} != {c['header']}")
            for r, lab, m in rows:
                p0, p1 = _month_period(year, m) if m else (f"{year}-01-01", f"{year}-12-31")
                value, printed, status, formula = cell(r, col)
                out.append(Obs(cfg["source_id"], no, c.get("variable", t["variable"]), c["level"], p0, p1,
                               "monthly" if m else "annual", value, printed, status, t["unit_printed"],
                               lab, " / ".join(c["header"]), "xlsx", plant_id=c["plant_id"], fuel=t.get("fuel"),
                               unit_group=c.get("unit_group"),
                               sheet=t["sheet"], cell_ref=f"{col}{r}", cell_formula=formula))
    return out
